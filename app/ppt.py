from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
def render(deck,path,template=None):
    prs=Presentation(template) if template else Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    for x in deck["slides"]:
        sl=prs.slides.add_slide(blank); title=sl.shapes.add_textbox(Inches(.7),Inches(.45),Inches(12),Inches(.7)); p=title.text_frame.paragraphs[0];p.text=x["title"];p.font.size=Pt(26);p.font.bold=True;p.font.color.rgb=RGBColor.from_string(deck["theme"]["primary"])
        box=sl.shapes.add_textbox(Inches(.8),Inches(1.45),Inches(11.7),Inches(5.3));tf=box.text_frame;c=x["content"]
        if c.get("subheadline"):tf.paragraphs[0].text=c["subheadline"]
        for b in c.get("bullets",[]):q=tf.add_paragraph();q.text=str(b);q.font.size=Pt(18);q.space_after=Pt(8)
        if c.get("key_takeaway"):q=tf.add_paragraph();q.text="Key takeaway: "+c["key_takeaway"];q.font.bold=True
    Path(path).parent.mkdir(parents=True,exist_ok=True);prs.save(path);return path
