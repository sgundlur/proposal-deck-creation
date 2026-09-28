from io import BytesIO
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from pptx import Presentation
def extract(data,filename):
    ext=Path(filename).suffix.lower()
    if ext in {".txt",".md"}: return data.decode("utf-8","replace")
    if ext==".pdf": return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages)
    if ext==".docx": return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs if p.text.strip())
    if ext==".pptx":
        out=[]; prs=Presentation(BytesIO(data))
        for i,sl in enumerate(prs.slides,1):
            out.append(f"[SLIDE {i}]"); out.extend(s.text for s in sl.shapes if hasattr(s,"text") and s.text.strip())
        return "\n".join(out)
    return ""
